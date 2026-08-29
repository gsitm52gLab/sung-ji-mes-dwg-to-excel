#!/usr/bin/env python3
"""emit_sheet1.py — 도면에서 부재 단위 명세(Sheet1)를 뽑아 구역별 엑셀로 만들고 대조한다.

발주 엑셀의 `Sheet1` 은 부재 한 줄씩의 명세이고, 그 AE/AF/AG 컬럼이
도면 `DECK_CTT` 의 (길이, 장수, 잔여) 와 같은 값이다. 즉 Sheet1 한 행 =
DECK_CTT 라벨 하나다. 이 대응을 이용해 도면에서 Sheet1 을 재구성한다.
"""

from __future__ import annotations

import glob
import os
import re
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook

import deckreport as R
import poc_deckorder as M
from deckconfig import cfg
from deckcheck.models import PIECE_HEADER

ZONE_RE = re.compile(r"(B1F-[^_.]+)")


def zone_of(path):
    """파일명에서 구역을 집는다.

    macOS 파일명은 NFD(자모 분해)라 '사' 가 'ᄉ+ᅡ' 로 저장된다. 도면 텍스트는
    NFC 이므로 정규화하지 않으면 한글 구역이 하나도 안 맞는다.
    """
    name = unicodedata.normalize("NFC", os.path.basename(path))
    m = ZONE_RE.search(name)
    return m.group(1) if m else os.path.splitext(name)[0][:20]


def has_sheet1(path):
    z = zipfile.ZipFile(path)
    names = re.findall(r'<sheet name="([^"]+)"',
                       z.read("xl/workbook.xml").decode("utf-8"))
    return "Sheet1" in names


def read_sheet1(path):
    """엑셀 Sheet1 → [(구간, 도면NO, SLAB, 길이, 장수, 잔여)]."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Sheet1"]
        rows = []
        for r in ws.iter_rows(values_only=True):
            if len(r) < 33 or not r[0]:
                continue
            zone, dcn, slab = r[0], r[1], r[2]
            ae, af, ag = r[30], r[31], r[32]   # AE, AF, AG
            if ae is None or af is None:
                continue
            try:
                rows.append((str(zone).strip(), dcn, str(slab).strip(),
                             int(float(ae)), int(float(af)),
                             int(float(ag)) if ag is not None else None))
            except (TypeError, ValueError):
                continue
        return rows
    finally:
        wb.close()


def roll_zone(zone):
    """도면의 3단계 구간 코드를 발주서의 2단계로 접는다.

    도면에는 '사-1-1' ~ '사-1-9', '사-2-1' ~ '사-2-5' 처럼 한 단계 더 쪼갠
    라벨이 있는데, 발주서는 '사-1' / '사-2' 까지만 쓴다.
    """
    parts = zone.split("-")
    return "-".join(parts[:2]) if len(parts) > 2 else zone


def drawing_rows(assigned, names, dcns):
    """배정 결과 → Sheet1 형태 행. 부재 하나가 한 행이다."""
    out = []
    for zone in sorted(assigned):
        for x, y, length, count, rem in sorted(assigned[zone], key=lambda p: (-p[1], p[0])):
            out.append([
                roll_zone(zone),
                M.nearest((x, y), dcns),
                M.nearest((x, y), names) or "",
                length, count, rem,
            ])
    return out


def write_zone(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(list(PIECE_HEADER))
    for r in rows:
        ws.append(r)
    wb.save(path)


def main():
    os.makedirs(cfg.by_zone_dir, exist_ok=True)
    polys, div, labels, pieces, names, dcns = M.extract_shop(cfg.shop_dxf)
    assigned = M.strat_divider_cells(polys, div, labels, pieces)
    rows = drawing_rows(assigned, names, dcns)
    R.console.print(
        f"도면 부재 {len(pieces)}개 → 배정 {len(rows)}행 / 구간 {len(assigned)}개")

    paths = [p for p in sorted(glob.glob(cfg.excel_glob))
             if not os.path.basename(p).startswith("~$")]

    R.heading("구역별 Sheet1 (부재 단위 명세)")
    t = R.table("구역", "엑셀부재", "도면부재", "부재일치", "구간일치", "파일", "비고")
    # 대조에서 빠진 사유는 길어서 표를 밀어낸다. 표에는 짧은 딱지만 두고
    # 사유 전문은 표 아래에 모아 찍는다.
    skipped = []
    grand = [0, 0, 0]
    for path in paths:
        zone = zone_of(path)
        prefix = zone.split("-")[-1]
        mine = [r for r in rows if r[0].split("-")[0] == prefix]
        out = os.path.join(cfg.by_zone_dir, f"Sheet1_{zone}.xlsx")
        write_zone(mine, out)
        name = os.path.basename(out)

        if not has_sheet1(path):
            t.add_row(zone, R.blank(), str(len(mine)), R.blank(), R.blank(),
                      name, R.note("Sheet1 없음"))
            skipped.append((zone, "엑셀에 Sheet1 시트가 없어 대조할 원본이 없다"))
            continue

        truth = read_sheet1(path)
        # 파일명 구역과 Sheet1 내용의 구간이 다르면 원본 워크북의 잔존 데이터다.
        # (실제로 B1F-나 파일의 Sheet1 은 B1F-사 것과 완전히 동일하다.)
        content = {t2[0].split("-")[0] for t2 in truth}
        if content and prefix not in content:
            t.add_row(zone, str(len(truth)), str(len(mine)), R.blank(), R.blank(),
                      name, R.note("대조 제외", style="yellow"))
            skipped.append((zone, f"Sheet1 내용이 {sorted(content)} 구역이다 "
                                  f"— 원본 워크북의 잔존 데이터"))
            continue
        # (길이, 장수) 로 짝지어 개수 기준 일치를 센다
        t_pair = Counter((t2[3], t2[4]) for t2 in truth)
        m_pair = Counter((r[3], r[4]) for r in mine)
        piece_hit = sum((t_pair & m_pair).values())
        # 구간까지 같은 것
        t_zone = Counter((t2[0], t2[3], t2[4]) for t2 in truth)
        m_zone = Counter((r[0], r[3], r[4]) for r in mine)
        zone_hit = sum((t_zone & m_zone).values())

        grand[0] += len(truth); grand[1] += piece_hit; grand[2] += zone_hit
        t.add_row(zone, str(len(truth)), str(len(mine)),
                  R.ratio(piece_hit, len(truth)), R.ratio(zone_hit, len(truth)),
                  name, "")
    R.console.print(t)

    if skipped:
        R.heading("대조 제외 사유")
        for zone, reason in skipped:
            R.detail(f"{zone}: {reason}", mark="⚠", style="yellow")

    if grand[0]:
        R.summary(f"대조 가능 {grand[0]}행 — 부재(길이·장수) 일치 {grand[1]} "
                  f"({grand[1]/grand[0]*100:.0f}%) / 구간까지 일치 {grand[2]} "
                  f"({grand[2]/grand[0]*100:.0f}%)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
