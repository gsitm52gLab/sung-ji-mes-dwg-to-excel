"""엑셀 `일람표` 시트 → DeckRecord."""

from __future__ import annotations

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import (
    CONSTANT_COLUMNS,
    EXCEL_HEADER,
    Cell,
    DeckRecord,
    HeaderMismatchError,
    SheetMissingError,
)
from .normalize import norm_camber, norm_lattice, norm_text

SHEET_NAME = "일람표"

# 컬럼별 정규화 규칙. 여기 없는 컬럼은 norm_text.
_NORMALIZERS = {"래티스": norm_lattice, "캠버": norm_camber}


def load_schedule(xlsx_path: str) -> list[DeckRecord]:
    """`.xlsm`/`.xlsx`의 `일람표` 시트를 레코드 목록으로 읽는다."""
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise SheetMissingError(
                f"{xlsx_path}: '{SHEET_NAME}' 시트가 없습니다"
            )
        rows = list(workbook[SHEET_NAME].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise HeaderMismatchError(f"{xlsx_path}: '{SHEET_NAME}' 시트가 비어 있습니다")

    width = len(EXCEL_HEADER)
    header = tuple(norm_text(v) for v in rows[0][:width])
    if header != EXCEL_HEADER:
        raise HeaderMismatchError(
            f"{xlsx_path}: '{SHEET_NAME}' 헤더 불일치\n"
            f"  기대: {list(EXCEL_HEADER)}\n"
            f"  실제: {list(header)}"
        )

    records: list[DeckRecord] = []
    for row_no, row in enumerate(rows[1:], start=2):
        values = [norm_text(v) for v in row[:width]]
        values += [""] * (width - len(values))
        symbol = values[0]
        if not symbol:
            continue

        def cell(column: str) -> Cell:
            index = EXCEL_HEADER.index(column)
            raw = values[index]
            normalizer = _NORMALIZERS.get(column, norm_text)
            return Cell(raw, normalizer(raw), f"{get_column_letter(index + 1)}{row_no}")

        fields = {
            column: cell(column)
            for column in EXCEL_HEADER[1:]  # SLAB NAME 은 키이므로 제외
        }
        records.append(DeckRecord(symbol=symbol, fields=fields, source=xlsx_path))

    return records
