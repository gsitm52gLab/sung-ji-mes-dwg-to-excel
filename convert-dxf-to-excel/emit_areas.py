#!/usr/bin/env python3
"""emit_areas.py — 도면 기준으로 구역별 발주 엑셀을 만든다.

고객 발주서는 구역 하나에 파일 하나이고, 그 안에 제작의뢰서 / Sheet1 / 일람표
시트가 들어 있다. 이 스크립트는 같은 규칙으로 도면에서 파일을 만든다.

앞선 스크립트들은 `excel/` 에 있는 발주서를 훑으며 그 구역만 처리했다.
그래서 발주서가 아직 없는 구역(가·다·라·아·자·카·1)은 산출되지 않았다.
여기서는 도면의 구간 라벨에서 구역 목록을 직접 뽑으므로 전 구역이 나온다.
발주서가 있는 구역은 대조 결과도 함께 출력한다.

사용:
    python3 emit_areas.py
"""

from __future__ import annotations

import glob
import os
import sys
import unicodedata
from collections import defaultdict

from openpyxl import Workbook

import deckreport as R
import poc_deckorder as M
from deckconfig import cfg
from emit_order import norm_slab, order_count, roll_zone
from deckcheck.models import EXCEL_HEADER, ORDER_HEADER, PIECE_HEADER

# 발주 대상이 아닌 라벨. 단열 구간과 메모성 텍스트는 제작 부재가 아니다.
SKIP_AREAS = {"지붕", "B1F"}


def area_of(zone):
    """구간 코드에서 구역을 집는다. '사-1-3' → '사', 'B-4' → 'B'."""
    return zone.split("-")[0]


def collect(assigned, names, dcns):
    """배정 결과를 구역 → 구간 → 부재 로 정리한다."""
    areas = defaultdict(list)
    for zone, ps in assigned.items():
        rolled = roll_zone(zone)
        area = area_of(rolled)
        if area in SKIP_AREAS or "-" not in rolled:
            continue
        for x, y, length, count, rem in ps:
            areas[area].append({
                "구간": rolled,
                "도면NO": M.nearest((x, y), dcns),
                "SLAB": norm_slab(M.nearest((x, y), names) or ""),
                "길이": length,
                "장수": count,
                "잔여": rem,
                "발주장수": order_count(count, rem),
            })
    return areas


def order_rows(pieces, master):
    """부재 → 제작의뢰서 행. (구간, SLAB, 길이) 로 합산한다."""
    agg = defaultdict(int)
    dcn = {}
    for p in pieces:
        key = (p["구간"], p["SLAB"], p["길이"])
        agg[key] += p["발주장수"]
        dcn.setdefault(key, p["도면NO"])

    rows = []
    for (zone, slab, length) in sorted(agg):
        m = master.get(slab) or {}
        typ = m.get("TYPE") or ""
        total = agg[(zone, slab, length)]
        rows.append([
            zone, dcn[(zone, slab, length)], slab,
            typ[1:] or None, m.get("TG"), m.get("하부피복"), m.get("캠버"),
            length, total,
            f"{typ}-{m.get('TG')}" if typ else None,
            round(total * length / 1000 * M.DECK_WIDTH_M, 3),
        ])
    return rows


def schedule_rows(pieces, master):
    """그 구역이 실제로 쓰는 데크 타입만 일람표로 낸다."""
    used = sorted({p["SLAB"] for p in pieces if p["SLAB"] in master})
    rows = []
    for slab in used:
        f = master[slab]
        rows.append([slab] + [
            None if col in ("단부재", "강판타입") else f.get(col)
            for col in EXCEL_HEADER[1:]
        ])
    return rows


def write_area(area, pieces, master, path):
    wb = Workbook()

    ws = wb.active
    ws.title = "제작의뢰서"
    ws.append(list(ORDER_HEADER))
    orders = order_rows(pieces, master)
    for r in orders:
        ws.append(r)

    s1 = wb.create_sheet("Sheet1")
    s1.append(list(PIECE_HEADER))
    for p in sorted(pieces, key=lambda p: (p["구간"], -p["길이"])):
        s1.append([p["구간"], p["도면NO"], p["SLAB"],
                   p["길이"], p["발주장수"], p["잔여"]])

    sc = wb.create_sheet("일람표")
    sc.append(list(EXCEL_HEADER))
    for r in schedule_rows(pieces, master):
        sc.append(r)

    wb.save(path)
    return len(orders)


def existing_orders():
    """구역 → 기존 발주서 경로. 파일명이 NFD 라 정규화해서 찾는다."""
    found = {}
    for p in sorted(glob.glob(cfg.excel_glob)):
        name = unicodedata.normalize("NFC", os.path.basename(p))
        if name.startswith("~$"):
            continue
        for part in name.replace(".xlsm", "").split("_"):
            if part.startswith("B1F-"):
                found[part[4:].split(".")[0]] = p
    return found


def main():
    os.makedirs(cfg.areas_dir, exist_ok=True)
    polys, div, labels, pieces, names, dcns = M.extract_shop(cfg.shop_dxf)
    master = M.load_type_master(cfg.detail_dxf)
    assigned = M.strat_divider_cells(polys, div, labels, pieces)
    areas = collect(assigned, names, dcns)

    oracle = M.load_oracle(cfg.excel_glob)
    truth = defaultdict(dict)
    for r in oracle:
        truth[area_of(r["구간"])][(r["구간"], r["SLAB"], r["길이"])] = \
            r["합계"] - r["강판"]

    have = existing_orders()
    R.console.print(f"도면 구역 {len(areas)}개 / 기존 발주서 {len(have)}개")

    R.heading("구역별 발주 엑셀")
    t = R.table("구역", "구간", "도면부재", "도면의뢰행", "장수", "대조", "파일")
    tot = hit = 0
    for area in sorted(areas, key=lambda a: -sum(p["발주장수"] for p in areas[a])):
        ps = areas[area]
        path = os.path.join(cfg.areas_dir, f"B1F-{area}.xlsx")
        n_rows = write_area(area, ps, master, path)
        qty = sum(p["발주장수"] for p in ps)
        zones = len({p["구간"] for p in ps})

        if area in truth:
            mine = {(r[0], r[2], r[7]): r[8] for r in order_rows(ps, master)}
            truth_area = truth[area]
            ok = sum(1 for k, v in truth_area.items() if mine.get(k) == v)
            tot += len(truth_area); hit += ok
            verdict = R.ratio(ok, len(truth_area), "행 일치")
        else:
            verdict = R.note("발주서 없음 — 신규", style="cyan")

        t.add_row(area, str(zones), str(len(ps)), str(n_rows), str(qty),
                  verdict, os.path.basename(path))
    R.console.print(t)

    R.summary(f"기존 발주서가 있는 구역: {hit}/{tot} 행 일치 ({hit/tot*100:.0f}%)")
    R.footnote(f"→ {cfg.areas_dir}/"
               f"  (구역당 파일 1개, 시트 3개: 제작의뢰서 / Sheet1 / 일람표)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
